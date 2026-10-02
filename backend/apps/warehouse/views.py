"""
仓库管理视图
"""
import logging
import io
from django.http import HttpResponse
from django.utils import timezone
from django.utils.dateparse import parse_datetime
from rest_framework.views import APIView
from rest_framework.permissions import IsAuthenticated
from rest_framework.parsers import MultiPartParser, FormParser, JSONParser
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from apps.core.response import success_response, error_response
from .models import (
    Unit, Category, Variety, Goods, StockIn, StockOut, Warning, Approval,
    StorageLocation, StorageReservation,
)
from .serializers import (
    UnitSerializer, UnitCreateSerializer,
    CategorySerializer, CategoryCreateSerializer,
    VarietySerializer, VarietyCreateSerializer,
    GoodsSerializer, StockInSerializer, StockInCreateSerializer, StockOutSerializer,
    WarningSerializer, ApprovalSerializer,
    StorageLocationSerializer, StorageLocationCreateSerializer,
    StorageReservationSerializer, StorageReservationCreateSerializer,
    StorageReservationReleaseSerializer, StorageReservationRescheduleSerializer,
)
from . import services

logger = logging.getLogger('apps')


# ==================== 单位管理 ====================

class UnitListView(APIView):
    """单位列表视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        queryset = Unit.objects.all().order_by('-created_at')
        
        page = int(request.query_params.get('page', 1))
        page_size = int(request.query_params.get('page_size', 10))
        start = (page - 1) * page_size
        end = start + page_size
        
        total = queryset.count()
        units = queryset[start:end]
        
        serializer = UnitSerializer(units, many=True)
        
        return success_response(data={
            'list': serializer.data,
            'total': total,
            'page': page,
            'page_size': page_size
        })
    
    def post(self, request):
        """创建单位"""
        serializer = UnitCreateSerializer(data=request.data)
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0][0]
            return error_response(message=str(first_error))
        
        unit = Unit.objects.create(
            name=serializer.validated_data['name'],
            created_by=request.user
        )
        
        logger.info(f"User {request.user.username} created unit {unit.name}")
        
        return success_response(data=UnitSerializer(unit).data, message='创建成功')


class UnitDetailView(APIView):
    """单位详情视图"""
    permission_classes = [IsAuthenticated]
    
    def put(self, request, pk):
        """更新单位"""
        try:
            unit = Unit.objects.get(pk=pk)
        except Unit.DoesNotExist:
            return error_response(message='单位不存在', code=404)
        
        serializer = UnitCreateSerializer(data=request.data, context={'instance': unit})
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0][0]
            return error_response(message=str(first_error))
        
        unit.name = serializer.validated_data['name']
        unit.save()
        
        logger.info(f"User {request.user.username} updated unit {unit.name}")
        
        return success_response(data=UnitSerializer(unit).data, message='更新成功')
    
    def delete(self, request, pk):
        """删除单位"""
        try:
            unit = Unit.objects.get(pk=pk)
        except Unit.DoesNotExist:
            return error_response(message='单位不存在', code=404)
        
        if unit.is_linked:
            return error_response(message='该单位已被关联，无法删除')
        
        name = unit.name
        unit.delete()
        
        logger.info(f"User {request.user.username} deleted unit {name}")
        
        return success_response(message='删除成功')


class UnitBatchDeleteView(APIView):
    """单位批量删除视图"""
    permission_classes = [IsAuthenticated]
    
    def post(self, request):
        ids = request.data.get('ids', [])
        if not ids:
            return error_response(message='请选择要删除的单位')
        
        # 只删除未关联的单位
        units = Unit.objects.filter(pk__in=ids)
        deleted_count = 0
        for unit in units:
            if not unit.is_linked:
                unit.delete()
                deleted_count += 1
        
        logger.info(f"User {request.user.username} batch deleted {deleted_count} units")
        
        return success_response(message=f'成功删除 {deleted_count} 个单位')


class UnitAllView(APIView):
    """获取所有单位（用于下拉选择）"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        units = Unit.objects.filter(is_active=True).order_by('name')
        serializer = UnitSerializer(units, many=True)
        return success_response(data=serializer.data)


# ==================== 品类管理 ====================

class CategoryListView(APIView):
    """品类列表视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        queryset = Category.objects.all().order_by('-created_at')
        
        page = int(request.query_params.get('page', 1))
        page_size = int(request.query_params.get('page_size', 10))
        start = (page - 1) * page_size
        end = start + page_size
        
        total = queryset.count()
        categories = queryset[start:end]
        
        serializer = CategorySerializer(categories, many=True)
        
        return success_response(data={
            'list': serializer.data,
            'total': total,
            'page': page,
            'page_size': page_size
        })
    
    def post(self, request):
        """创建品类"""
        serializer = CategoryCreateSerializer(data=request.data)
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0][0]
            return error_response(message=str(first_error))
        
        unit = Unit.objects.get(pk=serializer.validated_data['unit'])
        category = Category.objects.create(
            name=serializer.validated_data['name'],
            unit=unit,
            created_by=request.user
        )
        
        logger.info(f"User {request.user.username} created category {category.name}")
        
        return success_response(data=CategorySerializer(category).data, message='创建成功')


class CategoryDetailView(APIView):
    """品类详情视图"""
    permission_classes = [IsAuthenticated]
    
    def put(self, request, pk):
        """更新品类"""
        try:
            category = Category.objects.get(pk=pk)
        except Category.DoesNotExist:
            return error_response(message='品类不存在', code=404)
        
        serializer = CategoryCreateSerializer(data=request.data, context={'instance': category})
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0][0]
            return error_response(message=str(first_error))
        
        category.name = serializer.validated_data['name']
        category.unit = Unit.objects.get(pk=serializer.validated_data['unit'])
        category.save()
        
        logger.info(f"User {request.user.username} updated category {category.name}")
        
        return success_response(data=CategorySerializer(category).data, message='更新成功')
    
    def delete(self, request, pk):
        """删除品类"""
        try:
            category = Category.objects.get(pk=pk)
        except Category.DoesNotExist:
            return error_response(message='品类不存在', code=404)
        
        if category.is_linked:
            return error_response(message='该品类已被关联，无法删除')
        
        name = category.name
        category.delete()
        
        logger.info(f"User {request.user.username} deleted category {name}")
        
        return success_response(message='删除成功')


class CategoryBatchDeleteView(APIView):
    """品类批量删除视图"""
    permission_classes = [IsAuthenticated]
    
    def post(self, request):
        ids = request.data.get('ids', [])
        if not ids:
            return error_response(message='请选择要删除的品类')
        
        categories = Category.objects.filter(pk__in=ids)
        deleted_count = 0
        for category in categories:
            if not category.is_linked:
                category.delete()
                deleted_count += 1
        
        logger.info(f"User {request.user.username} batch deleted {deleted_count} categories")
        
        return success_response(message=f'成功删除 {deleted_count} 个品类')


class CategoryAllView(APIView):
    """获取所有品类（用于下拉选择）"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        categories = Category.objects.filter(is_active=True).order_by('name')
        serializer = CategorySerializer(categories, many=True)
        return success_response(data=serializer.data)


# ==================== 品种管理 ====================

class VarietyListView(APIView):
    """品种列表视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        queryset = Variety.objects.all().order_by('-created_at')
        
        page = int(request.query_params.get('page', 1))
        page_size = int(request.query_params.get('page_size', 10))
        start = (page - 1) * page_size
        end = start + page_size
        
        total = queryset.count()
        varieties = queryset[start:end]
        
        serializer = VarietySerializer(varieties, many=True)
        
        return success_response(data={
            'list': serializer.data,
            'total': total,
            'page': page,
            'page_size': page_size
        })
    
    def post(self, request):
        """创建品种"""
        serializer = VarietyCreateSerializer(data=request.data)
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0]
            if isinstance(first_error, list):
                first_error = first_error[0]
            return error_response(message=str(first_error))
        
        category = Category.objects.get(pk=serializer.validated_data['category'])
        variety = Variety.objects.create(
            name=serializer.validated_data['name'],
            category=category,
            created_by=request.user
        )
        
        logger.info(f"User {request.user.username} created variety {variety.name}")
        
        return success_response(data=VarietySerializer(variety).data, message='创建成功')


class VarietyDetailView(APIView):
    """品种详情视图"""
    permission_classes = [IsAuthenticated]
    
    def put(self, request, pk):
        """更新品种"""
        try:
            variety = Variety.objects.get(pk=pk)
        except Variety.DoesNotExist:
            return error_response(message='品种不存在', code=404)
        
        serializer = VarietyCreateSerializer(data=request.data, context={'instance': variety})
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0]
            if isinstance(first_error, list):
                first_error = first_error[0]
            return error_response(message=str(first_error))
        
        variety.name = serializer.validated_data['name']
        variety.category = Category.objects.get(pk=serializer.validated_data['category'])
        variety.save()
        
        logger.info(f"User {request.user.username} updated variety {variety.name}")
        
        return success_response(data=VarietySerializer(variety).data, message='更新成功')
    
    def delete(self, request, pk):
        """删除品种"""
        try:
            variety = Variety.objects.get(pk=pk)
        except Variety.DoesNotExist:
            return error_response(message='品种不存在', code=404)
        
        if variety.is_in_stock:
            return error_response(message='该品种已入库，无法删除')
        
        name = variety.name
        variety.delete()
        
        logger.info(f"User {request.user.username} deleted variety {name}")
        
        return success_response(message='删除成功')


class VarietyBatchDeleteView(APIView):
    """品种批量删除视图"""
    permission_classes = [IsAuthenticated]
    
    def post(self, request):
        ids = request.data.get('ids', [])
        if not ids:
            return error_response(message='请选择要删除的品种')
        
        varieties = Variety.objects.filter(pk__in=ids)
        deleted_count = 0
        for variety in varieties:
            if not variety.is_in_stock:
                variety.delete()
                deleted_count += 1
        
        logger.info(f"User {request.user.username} batch deleted {deleted_count} varieties")
        
        return success_response(message=f'成功删除 {deleted_count} 个品种')


class VarietyTemplateView(APIView):
    """品种导入模板下载"""
    permission_classes = []  # 允许匿名访问，通过token参数验证
    
    def get(self, request):
        # 从URL参数获取token进行验证
        from apps.authentication.backends import decode_token
        from apps.authentication.models import User
        
        token = request.query_params.get('token')
        if not token:
            return error_response(message='缺少认证信息', code=401)
        
        payload = decode_token(token)
        if not payload:
            return error_response(message='认证信息无效或已过期', code=401)
        
        try:
            user = User.objects.get(pk=payload['user_id'])
        except User.DoesNotExist:
            return error_response(message='用户不存在', code=401)
        
        wb = Workbook()
        
        # 第一个表格 - 导入模板
        ws1 = wb.active
        ws1.title = '品种导入'
        
        # 设置表头样式
        header_font = Font(bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='4F46E5', end_color='4F46E5', fill_type='solid')
        header_alignment = Alignment(horizontal='center', vertical='center')
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        headers = ['品种', '品类', '单位']
        for col, header in enumerate(headers, 1):
            cell = ws1.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border
        
        # 设置列宽
        ws1.column_dimensions['A'].width = 25
        ws1.column_dimensions['B'].width = 20
        ws1.column_dimensions['C'].width = 15
        
        # 第二个表格 - 品类参考
        ws2 = wb.create_sheet(title='品类参考')
        
        headers2 = ['品类', '单位']
        for col, header in enumerate(headers2, 1):
            cell = ws2.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border
        
        # 填充品类数据
        categories = Category.objects.filter(is_active=True).select_related('unit')
        for row, category in enumerate(categories, 2):
            ws2.cell(row=row, column=1, value=category.name).border = thin_border
            ws2.cell(row=row, column=2, value=category.unit.name).border = thin_border
        
        ws2.column_dimensions['A'].width = 20
        ws2.column_dimensions['B'].width = 15
        
        # 返回Excel文件
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        
        response = HttpResponse(
            output.read(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename=variety_import_template.xlsx'
        
        return response


class VarietyImportView(APIView):
    """品种导入视图"""
    permission_classes = [IsAuthenticated]
    parser_classes = [MultiPartParser, FormParser]
    
    def post(self, request):
        if 'file' not in request.FILES:
            return error_response(message='请上传文件')
        
        file = request.FILES['file']
        
        try:
            wb = load_workbook(file)
            ws = wb.active
        except Exception as e:
            return error_response(message='文件格式错误，请上传Excel文件')
        
        # 获取所有品类及其单位
        categories = {c.name: c for c in Category.objects.filter(is_active=True).select_related('unit')}
        
        can_import = []
        cannot_import = []
        
        for row in range(2, ws.max_row + 1):
            variety_name = ws.cell(row=row, column=1).value
            category_name = ws.cell(row=row, column=2).value
            unit_name = ws.cell(row=row, column=3).value
            
            if not variety_name:
                continue
            
            variety_name = str(variety_name).strip()
            category_name = str(category_name).strip() if category_name else ''
            unit_name = str(unit_name).strip() if unit_name else ''
            
            # 验证
            error_msg = None
            
            if not variety_name:
                error_msg = '品种名称不能为空'
            elif len(variety_name) > 20:
                error_msg = '品种名称最多20个字'
            elif not category_name:
                error_msg = '品类不能为空'
            elif category_name not in categories:
                error_msg = f'品类"{category_name}"不存在'
            elif not unit_name:
                error_msg = '单位不能为空'
            elif categories.get(category_name) and categories[category_name].unit.name != unit_name:
                error_msg = f'单位与品类不匹配，应为"{categories[category_name].unit.name}"'
            elif Variety.objects.filter(name=variety_name, category__name=category_name).exists():
                error_msg = '该品种已存在'
            
            if error_msg:
                cannot_import.append({
                    'row': row,
                    'variety': variety_name,
                    'category': category_name,
                    'unit': unit_name,
                    'reason': error_msg
                })
            else:
                can_import.append({
                    'row': row,
                    'variety': variety_name,
                    'category': category_name,
                    'unit': unit_name
                })
        
        # 如果是预览请求
        if request.data.get('preview') == 'true':
            return success_response(data={
                'can_import': can_import,
                'cannot_import': cannot_import,
                'can_import_count': len(can_import),
                'cannot_import_count': len(cannot_import)
            })
        
        # 执行导入
        imported_count = 0
        for item in can_import:
            category = categories[item['category']]
            Variety.objects.create(
                name=item['variety'],
                category=category,
                created_by=request.user
            )
            imported_count += 1
        
        logger.info(f"User {request.user.username} imported {imported_count} varieties")
        
        return success_response(
            data={
                'imported_count': imported_count,
                'failed_count': len(cannot_import),
                'failed_items': cannot_import
            },
            message=f'成功导入 {imported_count} 个品种'
        )


# ==================== 其他视图占位 ====================

class DashboardView(APIView):
    """仪表盘视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        return success_response(data={
            'message': '仪表盘功能开发中...'
        })


class GoodsListView(APIView):
    """货物列表视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        return success_response(data={
            'list': [],
            'total': 0,
            'page': 1,
            'page_size': 10
        })


class StockInListView(APIView):
    """入库记录列表视图"""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        queryset = StockIn.objects.select_related(
            'goods', 'operator', 'reservation__location'
        ).all().order_by('-stock_in_time')

        page = int(request.query_params.get('page', 1))
        page_size = int(request.query_params.get('page_size', 10))
        start = (page - 1) * page_size
        end = start + page_size

        total = queryset.count()
        records = queryset[start:end]
        serializer = StockInSerializer(records, many=True)

        return success_response(data={
            'list': serializer.data,
            'total': total,
            'page': page,
            'page_size': page_size
        })

    def post(self, request):
        """正式入库：必须消费一个有效预约"""
        serializer = StockInCreateSerializer(data=request.data)
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0]
            if isinstance(first_error, list):
                first_error = first_error[0]
            return error_response(message=str(first_error))

        return _occupy_and_respond(
            request,
            reservation_id=serializer.validated_data['reservation_id'],
            serializer=serializer,
        )


# ==================== 保管位管理 ====================

def _parse_datetime_param(value, name):
    """解析查询参数中的时间，支持ISO 8601，无时区时按当前时区处理"""
    parsed = parse_datetime(value)
    if parsed is None:
        raise ValueError(f'{name}格式错误，应为ISO 8601时间格式')
    if timezone.is_naive(parsed):
        parsed = timezone.make_aware(parsed, timezone.get_current_timezone())
    return parsed


def _occupy_and_respond(request, reservation_id, serializer):
    """执行正式入库（消费预约）并组装响应"""
    try:
        stock_in = services.occupy_reservation(
            reservation_id=reservation_id,
            goods=Goods.objects.get(pk=serializer.validated_data['goods']),
            operator=request.user,
            quantity=serializer.validated_data['quantity'],
            batch_no=serializer.validated_data.get('batch_no', ''),
            supplier=serializer.validated_data.get('supplier', ''),
            remark=serializer.validated_data.get('remark', ''),
        )
    except StorageReservation.DoesNotExist:
        return error_response(message='预约不存在', code=404)
    except services.ReservationStateError as e:
        return error_response(message=e.message)

    logger.info(f"User {request.user.username} stocked in via reservation {reservation_id}")
    return success_response(data=StockInSerializer(stock_in).data, message='入库成功')


class StorageLocationListView(APIView):
    """保管位列表视图"""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        queryset = StorageLocation.objects.all().order_by('name')

        page = int(request.query_params.get('page', 1))
        page_size = int(request.query_params.get('page_size', 10))
        start = (page - 1) * page_size
        end = start + page_size

        total = queryset.count()
        locations = queryset[start:end]
        serializer = StorageLocationSerializer(locations, many=True)

        return success_response(data={
            'list': serializer.data,
            'total': total,
            'page': page,
            'page_size': page_size
        })

    def post(self, request):
        """创建保管位"""
        serializer = StorageLocationCreateSerializer(data=request.data)
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0][0]
            return error_response(message=str(first_error))

        location = StorageLocation.objects.create(
            name=serializer.validated_data['name'],
            capacity=serializer.validated_data['capacity'],
            description=serializer.validated_data.get('description', ''),
            is_active=serializer.validated_data.get('is_active', True),
            created_by=request.user
        )

        logger.info(f"User {request.user.username} created storage location {location.name}")

        return success_response(data=StorageLocationSerializer(location).data, message='创建成功')


class StorageLocationDetailView(APIView):
    """保管位详情视图"""
    permission_classes = [IsAuthenticated]

    def put(self, request, pk):
        """更新保管位"""
        try:
            location = StorageLocation.objects.get(pk=pk)
        except StorageLocation.DoesNotExist:
            return error_response(message='保管位不存在', code=404)

        serializer = StorageLocationCreateSerializer(
            data=request.data, context={'instance': location}
        )
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0][0]
            return error_response(message=str(first_error))

        new_capacity = serializer.validated_data['capacity']
        holding = location.reservations.filter(
            status__in=services.HOLDING_STATUSES
        ).exists()
        if holding and new_capacity < location.capacity:
            return error_response(message='该保管位存在有效预约或占用，不能缩减容量')

        location.name = serializer.validated_data['name']
        location.capacity = new_capacity
        location.description = serializer.validated_data.get('description', '')
        location.is_active = serializer.validated_data.get('is_active', True)
        location.save()

        logger.info(f"User {request.user.username} updated storage location {location.name}")

        return success_response(data=StorageLocationSerializer(location).data, message='更新成功')

    def delete(self, request, pk):
        """删除保管位（存在预约或占用记录的保管位不可删除，应停用）"""
        try:
            location = StorageLocation.objects.get(pk=pk)
        except StorageLocation.DoesNotExist:
            return error_response(message='保管位不存在', code=404)

        if location.reservations.filter(status__in=services.HOLDING_STATUSES).exists():
            return error_response(message='该保管位存在有效预约或占用，无法删除')
        if location.reservations.exists():
            return error_response(message='该保管位已产生预约记录，为保证可追溯不可删除，可将其停用')

        name = location.name
        location.delete()

        logger.info(f"User {request.user.username} deleted storage location {name}")

        return success_response(message='删除成功')


class StorageLocationAvailabilityView(APIView):
    """保管位可用性查询：返回冲突区间与剩余容量"""
    permission_classes = [IsAuthenticated]

    def get(self, request, pk):
        if not StorageLocation.objects.filter(pk=pk).exists():
            return error_response(message='保管位不存在', code=404)

        start_param = request.query_params.get('start')
        end_param = request.query_params.get('end')
        if not start_param or not end_param:
            return error_response(message='请提供start和end查询参数')

        try:
            start = _parse_datetime_param(start_param, 'start')
            end = _parse_datetime_param(end_param, 'end')
        except ValueError as e:
            return error_response(message=str(e))

        if start >= end:
            return error_response(message='结束时间必须晚于开始时间')

        try:
            quantity = int(request.query_params.get('quantity', 1))
        except (TypeError, ValueError):
            return error_response(message='quantity必须为整数')
        if quantity < 1:
            return error_response(message='quantity必须大于0')

        availability = services.get_availability(pk, start, end, quantity=quantity)
        return success_response(data=availability)


# ==================== 保管位预约管理 ====================

class StorageReservationListView(APIView):
    """保管位预约列表视图"""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        # 懒过期，保证值班员看到最新状态
        services.expire_due_reservations()

        queryset = StorageReservation.objects.select_related(
            'location', 'applicant'
        ).all().order_by('-created_at')

        location_id = request.query_params.get('location')
        if location_id:
            queryset = queryset.filter(location_id=location_id)
        status_param = request.query_params.get('status')
        if status_param:
            queryset = queryset.filter(status=status_param)
        team_name = request.query_params.get('team_name')
        if team_name:
            queryset = queryset.filter(team_name__icontains=team_name)

        # 时段重叠过滤：start_time < end 且 end_time > start
        overlap_start = request.query_params.get('start')
        overlap_end = request.query_params.get('end')
        try:
            if overlap_start and overlap_end:
                start = _parse_datetime_param(overlap_start, 'start')
                end = _parse_datetime_param(overlap_end, 'end')
                queryset = queryset.filter(start_time__lt=end, end_time__gt=start)
        except ValueError as e:
            return error_response(message=str(e))

        page = int(request.query_params.get('page', 1))
        page_size = int(request.query_params.get('page_size', 10))
        start_idx = (page - 1) * page_size
        end_idx = start_idx + page_size

        total = queryset.count()
        reservations = queryset[start_idx:end_idx]
        serializer = StorageReservationSerializer(reservations, many=True)

        return success_response(data={
            'list': serializer.data,
            'total': total,
            'page': page,
            'page_size': page_size
        })

    def post(self, request):
        """创建预约：容量不足时返回冲突区间与剩余容量"""
        serializer = StorageReservationCreateSerializer(data=request.data)
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0]
            if isinstance(first_error, list):
                first_error = first_error[0]
            return error_response(message=str(first_error))

        location = StorageLocation.objects.get(pk=serializer.validated_data['location'])
        try:
            reservation = services.create_reservation(
                location=location,
                applicant=request.user,
                team_name=serializer.validated_data['team_name'],
                quantity=serializer.validated_data['quantity'],
                start=serializer.validated_data['start_time'],
                end=serializer.validated_data['end_time'],
                remark=serializer.validated_data.get('remark', ''),
            )
        except services.CapacityConflictError as e:
            return error_response(message=e.message, code=409, data=e.detail)
        except services.ReservationStateError as e:
            return error_response(message=e.message)

        logger.info(
            f"User {request.user.username} reserved location {location.name} "
            f"for team {reservation.team_name}"
        )

        return success_response(
            data=StorageReservationSerializer(reservation).data, message='预约成功'
        )


class StorageReservationDetailView(APIView):
    """保管位预约详情视图"""
    permission_classes = [IsAuthenticated]

    def get(self, request, pk):
        services.expire_due_reservations()
        try:
            reservation = StorageReservation.objects.select_related(
                'location', 'applicant'
            ).get(pk=pk)
        except StorageReservation.DoesNotExist:
            return error_response(message='预约不存在', code=404)

        return success_response(data=StorageReservationSerializer(reservation).data)


class StorageReservationReleaseView(APIView):
    """预约释放视图：支持整体释放与部分取消"""
    permission_classes = [IsAuthenticated]

    def post(self, request, pk):
        serializer = StorageReservationReleaseSerializer(data=request.data)
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0]
            if isinstance(first_error, list):
                first_error = first_error[0]
            return error_response(message=str(first_error))

        try:
            reservation = services.release_reservation(
                pk, quantity=serializer.validated_data.get('quantity')
            )
        except StorageReservation.DoesNotExist:
            return error_response(message='预约不存在', code=404)
        except services.ReservationStateError as e:
            return error_response(message=e.message)

        logger.info(f"User {request.user.username} released reservation {pk}")

        return success_response(
            data=StorageReservationSerializer(reservation).data, message='释放成功'
        )


class StorageReservationRescheduleView(APIView):
    """预约改期视图：先验证新时段，通过后才释放旧时段"""
    permission_classes = [IsAuthenticated]

    def post(self, request, pk):
        serializer = StorageReservationRescheduleSerializer(data=request.data)
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0]
            if isinstance(first_error, list):
                first_error = first_error[0]
            return error_response(message=str(first_error))

        try:
            reservation = services.reschedule_reservation(
                pk,
                start=serializer.validated_data['start_time'],
                end=serializer.validated_data['end_time'],
                quantity=serializer.validated_data.get('quantity'),
            )
        except StorageReservation.DoesNotExist:
            return error_response(message='预约不存在', code=404)
        except services.CapacityConflictError as e:
            return error_response(message=e.message, code=409, data=e.detail)
        except services.ReservationStateError as e:
            return error_response(message=e.message)

        logger.info(f"User {request.user.username} rescheduled reservation {pk}")

        return success_response(
            data=StorageReservationSerializer(reservation).data, message='改期成功'
        )


class StorageReservationOccupyView(APIView):
    """预约占用视图：正式入库，消费有效预约"""
    permission_classes = [IsAuthenticated]

    def post(self, request, pk):
        data = {**request.data, 'reservation_id': pk}
        serializer = StockInCreateSerializer(data=data)
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0]
            if isinstance(first_error, list):
                first_error = first_error[0]
            return error_response(message=str(first_error))

        return _occupy_and_respond(request, reservation_id=pk, serializer=serializer)


class StockOutListView(APIView):
    """出库记录列表视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        return success_response(data={
            'list': [],
            'total': 0,
            'page': 1,
            'page_size': 10
        })


class WarningListView(APIView):
    """预警记录列表视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        return success_response(data={
            'list': [],
            'total': 0,
            'page': 1,
            'page_size': 10
        })


class ApprovalListView(APIView):
    """审批记录列表视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        return success_response(data={
            'list': [],
            'total': 0,
            'page': 1,
            'page_size': 10
        })
